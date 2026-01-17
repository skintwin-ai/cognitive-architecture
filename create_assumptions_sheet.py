import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os

# Define file path
excel_file = '/home/ubuntu/skin_care_financial_model.xlsx'

# --- Create Assumptions Sheet Data --- 
assumptions_data = {
    'Category': [
        # Global Assumptions
        'Global Assumptions', 'Global Assumptions', 'Global Assumptions', 'Global Assumptions', 'Global Assumptions', 'Global Assumptions',
        # Revenue Assumptions
        'Revenue Assumptions', 'Revenue Assumptions', 'Revenue Assumptions', 'Revenue Assumptions', 'Revenue Assumptions', 'Revenue Assumptions', 'Revenue Assumptions',
        # COGS Assumptions
        'COGS Assumptions', 'COGS Assumptions', 'COGS Assumptions', 'COGS Assumptions', 'COGS Assumptions', 'COGS Assumptions', 'COGS Assumptions',
        # Operational Cost Assumptions
        'Operational Cost Assumptions', 'Operational Cost Assumptions', 'Operational Cost Assumptions', 'Operational Cost Assumptions', 'Operational Cost Assumptions', 'Operational Cost Assumptions',
        # Packaging & Logistics Assumptions
        'Packaging & Logistics Assumptions', 'Packaging & Logistics Assumptions', 'Packaging & Logistics Assumptions', 'Packaging & Logistics Assumptions', 'Packaging & Logistics Assumptions', 'Packaging & Logistics Assumptions', 'Packaging & Logistics Assumptions',
        # SG&A Assumptions
        'SG&A Assumptions', 'SG&A Assumptions', 'SG&A Assumptions', 'SG&A Assumptions', 'SG&A Assumptions', 'SG&A Assumptions', 'SG&A Assumptions', 'SG&A Assumptions',
        # R&D Assumptions
        'R&D Assumptions', 'R&D Assumptions', 'R&D Assumptions', 'R&D Assumptions', 'R&D Assumptions', 'R&D Assumptions',
        # HR Assumptions
        'HR Assumptions', 'HR Assumptions', 'HR Assumptions', 'HR Assumptions', 'HR Assumptions', 'HR Assumptions',
        # Capex Assumptions
        'Capex Assumptions', 'Capex Assumptions', 'Capex Assumptions', 'Capex Assumptions', 'Capex Assumptions', 'Capex Assumptions',
        # Working Capital Assumptions
        'Working Capital Assumptions', 'Working Capital Assumptions', 'Working Capital Assumptions', 'Working Capital Assumptions',
        # Group Assumptions
        'Group Assumptions', 'Group Assumptions', 'Group Assumptions', 'Group Assumptions', 'Group Assumptions',
        # Financial Assumptions
        'Financial Assumptions', 'Financial Assumptions', 'Financial Assumptions',
        # Sustainability & Compliance Assumptions
        'Sustainability & Compliance Assumptions', 'Sustainability & Compliance Assumptions', 'Sustainability & Compliance Assumptions', 'Sustainability & Compliance Assumptions', 'Sustainability & Compliance Assumptions',
    ],
    'Assumption': [
        # Global Assumptions
        'Projection Start Year', 'Projection Period (Years)', 'Currency', 'Inflation Rate (%)', 'Interest Rate (Base, %)', 'Corporate Tax Rate (%)',
        # Revenue Assumptions
        'Overall Market Growth Rate (%)', 'Product Category 1 Growth Rate (%)', 'Product Category 2 Growth Rate (%)', 'Channel 1 (DTC) Growth Rate (%)', 'Channel 2 (Retail) Growth Rate (%)', 'Avg. Price Increase per Year (%)', 'Seasonality Factor (Q1, Q2, Q3, Q4)',
        # COGS Assumptions
        'Raw Material Cost Increase (%)', 'Avg. Ingredient Cost per Unit (Base)', 'Manufacturing Labor Cost Increase (%)', 'Manufacturing Overhead Rate (% of Direct Labor)', 'Contract Manufacturing Fee (% of Revenue)', 'Quality Control Cost (% of COGS)', 'Yield Loss Rate (%)',
        # Operational Cost Assumptions
        'Facility Rent Increase (%)', 'Utilities Cost Increase (%)', 'Equipment Maintenance (% of Asset Value)', 'Inventory Holding Cost (% of Inventory Value)', 'Supply Chain Software Cost (Annual)', 'Sustainability Program Cost (Annual)',
        # Packaging & Logistics Assumptions
        'Packaging Material Cost Increase (%)', 'Warehousing Cost per Sq Ft (Annual)', 'Avg. Transportation Cost per Unit', 'Order Fulfillment Cost per Order', 'Last-Mile Delivery Cost (% of Revenue)', 'Returns Rate (% of Sales)', 'Logistics Software Cost (Annual)',
        # SG&A Assumptions
        'Sales Team Size Growth (%)', 'Avg. Sales Commission Rate (%)', 'Marketing Budget (% of Revenue)', 'Admin Staff Headcount Growth (%)', 'Office Rent Increase (%)', 'IT Support Cost (Per Employee)', 'Legal & Accounting Fees (Annual)', 'Travel & Entertainment (% of Revenue)',
        # R&D Assumptions
        'R&D Budget (% of Revenue)', 'Avg. Cost per New Product Launch', 'Clinical Trial Cost (Per Trial)', 'Patent Filing Cost (Per Patent)', 'Regulatory Compliance Cost (Annual)', 'Collaboration/Partnership Costs (Annual)',
        # HR Assumptions
        'Avg. Salary Increase (%)', 'Employee Benefits Cost (% of Salary)', 'Training Budget (Per Employee)', 'Recruitment Cost (Per Hire)', 'Employee Turnover Rate (%)', 'Temporary Staffing Cost (% of Payroll)',
        # Capex Assumptions
        'Annual Capex Budget (% of Revenue)', 'Avg. Equipment Useful Life (Years)', 'Facility Improvement Budget (Annual)', 'Technology Upgrade Cycle (Years)', 'Depreciation Method', 'Maintenance Capex (% of Total Capex)',
        # Working Capital Assumptions
        'Inventory Days - Raw Materials', 'Inventory Days - Finished Goods', 'Accounts Receivable Days', 'Accounts Payable Days',
        # Group Assumptions
        'Intercompany Sales Markup (%)', 'Transfer Pricing Policy', 'Shared Services Allocation Basis', 'Corporate Overhead Allocation (% of Revenue)', 'Consolidation Adjustments (Manual Input)',
        # Financial Assumptions
        'Debt Interest Rate (%)', 'Foreign Exchange Rate (Specify Currency)', 'Banking Fees (Annual)',
        # Sustainability & Compliance Assumptions
        'Certification Costs (Annual)', 'Environmental Compliance Fines (Estimate)', 'Audit Fees (Annual)', 'CSR Program Budget (Annual)', 'Supply Chain Audit Costs (Annual)',
    ],
    'Year 1': [
        # Global Assumptions
        '2025', '5', 'USD', '2.5%', '4.0%', '25.0%',
        # Revenue Assumptions
        '5.0%', '6.0%', '4.0%', '8.0%', '3.0%', '2.0%', '0.2, 0.3, 0.3, 0.2',
        # COGS Assumptions
        '3.0%', '$2.50', '3.5%', '20.0%', '10.0%', '2.0%', '3.0%',
        # Operational Cost Assumptions
        '3.0%', '4.0%', '2.5%', '5.0%', '$50,000', '$75,000',
        # Packaging & Logistics Assumptions
        '2.5%', '$15', '$0.75', '$3.50', '2.0%', '5.0%', '$35,000',
        # SG&A Assumptions
        '5.0%', '3.0%', '12.0%', '3.0%', '3.0%', '$1,200', '$120,000', '1.5%',
        # R&D Assumptions
        '5.0%', '$150,000', '$75,000', '$10,000', '$50,000', '$100,000',
        # HR Assumptions
        '3.5%', '25.0%', '$1,500', '$5,000', '15.0%', '5.0%',
        # Capex Assumptions
        '4.0%', '7', '$200,000', '3', 'Straight-Line', '30.0%',
        # Working Capital Assumptions
        '45', '30', '45', '30',
        # Group Assumptions
        '5.0%', 'Cost-Plus', 'Revenue-Based', '3.0%', '$0',
        # Financial Assumptions
        '5.5%', '1.00', '$25,000',
        # Sustainability & Compliance Assumptions
        '$30,000', '$10,000', '$45,000', '$60,000', '$25,000',
    ],
    'Year 2': [
        # Global Assumptions
        '2026', '5', 'USD', '2.3%', '4.0%', '25.0%',
        # Revenue Assumptions
        '5.2%', '6.5%', '4.2%', '9.0%', '3.2%', '2.2%', '0.2, 0.3, 0.3, 0.2',
        # COGS Assumptions
        '2.8%', '$2.55', '3.3%', '20.0%', '9.5%', '1.9%', '2.8%',
        # Operational Cost Assumptions
        '3.0%', '3.8%', '2.3%', '4.8%', '$51,250', '$80,000',
        # Packaging & Logistics Assumptions
        '2.4%', '$15.45', '$0.73', '$3.40', '1.9%', '4.8%', '$36,000',
        # SG&A Assumptions
        '4.5%', '3.0%', '11.5%', '2.8%', '3.0%', '$1,230', '$123,000', '1.4%',
        # R&D Assumptions
        '5.2%', '$155,000', '$77,000', '$10,200', '$51,250', '$105,000',
        # HR Assumptions
        '3.4%', '25.0%', '$1,550', '$5,100', '14.0%', '4.8%',
        # Capex Assumptions
        '3.8%', '7', '$205,000', '3', 'Straight-Line', '28.0%',
        # Working Capital Assumptions
        '43', '28', '43', '32',
        # Group Assumptions
        '5.0%', 'Cost-Plus', 'Revenue-Based', '2.9%', '$0',
        # Financial Assumptions
        '5.4%', '1.00', '$25,500',
        # Sustainability & Compliance Assumptions
        '$30,750', '$9,500', '$46,000', '$65,000', '$25,500',
    ],
    'Year 3': [
        # Global Assumptions
        '2027', '5', 'USD', '2.2%', '4.0%', '25.0%',
        # Revenue Assumptions
        '5.4%', '7.0%', '4.5%', '10.0%', '3.5%', '2.3%', '0.2, 0.3, 0.3, 0.2',
        # COGS Assumptions
        '2.7%', '$2.60', '3.2%', '20.0%', '9.0%', '1.8%', '2.6%',
        # Operational Cost Assumptions
        '3.0%', '3.6%', '2.2%', '4.6%', '$52,500', '$85,000',
        # Packaging & Logistics Assumptions
        '2.3%', '$15.90', '$0.71', '$3.30', '1.8%', '4.6%', '$37,000',
        # SG&A Assumptions
        '4.0%', '3.0%', '11.0%', '2.6%', '3.0%', '$1,260', '$126,000', '1.3%',
        # R&D Assumptions
        '5.4%', '$160,000', '$79,000', '$10,400', '$52,500', '$110,000',
        # HR Assumptions
        '3.3%', '25.0%', '$1,600', '$5,200', '13.0%', '4.6%',
        # Capex Assumptions
        '3.6%', '7', '$210,000', '3', 'Straight-Line', '26.0%',
        # Working Capital Assumptions
        '40', '26', '40', '34',
        # Group Assumptions
        '5.0%', 'Cost-Plus', 'Revenue-Based', '2.8%', '$0',
        # Financial Assumptions
        '5.3%', '1.00', '$26,000',
        # Sustainability & Compliance Assumptions
        '$31,500', '$9,000', '$47,000', '$70,000', '$26,000',
    ],
    'Year 4': [
        # Global Assumptions
        '2028', '5', 'USD', '2.1%', '4.0%', '25.0%',
        # Revenue Assumptions
        '5.6%', '7.5%', '4.8%', '11.0%', '3.8%', '2.4%', '0.2, 0.3, 0.3, 0.2',
        # COGS Assumptions
        '2.6%', '$2.65', '3.1%', '20.0%', '8.5%', '1.7%', '2.4%',
        # Operational Cost Assumptions
        '3.0%', '3.4%', '2.1%', '4.4%', '$53,750', '$90,000',
        # Packaging & Logistics Assumptions
        '2.2%', '$16.35', '$0.69', '$3.20', '1.7%', '4.4%', '$38,000',
        # SG&A Assumptions
        '3.5%', '3.0%', '10.5%', '2.4%', '3.0%', '$1,290', '$129,000', '1.2%',
        # R&D Assumptions
        '5.6%', '$165,000', '$81,000', '$10,600', '$53,750', '$115,000',
        # HR Assumptions
        '3.2%', '25.0%', '$1,650', '$5,300', '12.0%', '4.4%',
        # Capex Assumptions
        '3.4%', '7', '$215,000', '3', 'Straight-Line', '24.0%',
        # Working Capital Assumptions
        '38', '24', '38', '36',
        # Group Assumptions
        '5.0%', 'Cost-Plus', 'Revenue-Based', '2.7%', '$0',
        # Financial Assumptions
        '5.2%', '1.00', '$26,500',
        # Sustainability & Compliance Assumptions
        '$32,250', '$8,500', '$48,000', '$75,000', '$26,500',
    ],
    'Year 5': [
        # Global Assumptions
        '2029', '5', 'USD', '2.0%', '4.0%', '25.0%',
        # Revenue Assumptions
        '5.8%', '8.0%', '5.0%', '12.0%', '4.0%', '2.5%', '0.2, 0.3, 0.3, 0.2',
        # COGS Assumptions
        '2.5%', '$2.70', '3.0%', '20.0%', '8.0%', '1.6%', '2.2%',
        # Operational Cost Assumptions
        '3.0%', '3.2%', '2.0%', '4.2%', '$55,000', '$95,000',
        # Packaging & Logistics Assumptions
        '2.1%', '$16.80', '$0.67', '$3.10', '1.6%', '4.2%', '$39,000',
        # SG&A Assumptions
        '3.0%', '3.0%', '10.0%', '2.2%', '3.0%', '$1,320', '$132,000', '1.1%',
        # R&D Assumptions
        '5.8%', '$170,000', '$83,000', '$10,800', '$55,000', '$120,000',
        # HR Assumptions
        '3.1%', '25.0%', '$1,700', '$5,400', '11.0%', '4.2%',
        # Capex Assumptions
        '3.2%', '7', '$220,000', '3', 'Straight-Line', '22.0%',
        # Working Capital Assumptions
        '35', '22', '35', '38',
        # Group Assumptions
        '5.0%', 'Cost-Plus', 'Revenue-Based', '2.6%', '$0',
        # Financial Assumptions
        '5.1%', '1.00', '$27,000',
        # Sustainability & Compliance Assumptions
        '$33,000', '$8,000', '$49,000', '$80,000', '$27,000',
    ],
    'Notes/Drivers': [
        # Global Assumptions
        'Starting year for the financial projections.', 'Duration of the forecast.', 'Reporting currency (e.g., USD, EUR).', 'Assumed general inflation rate.', 'Base interest rate for debt calculations.', 'Effective corporate income tax rate.',
        # Revenue Assumptions
        'Assumed annual growth rate for the overall skin care market.', 'Specific growth rate for Product Category 1.', 'Specific growth rate for Product Category 2.', 'Growth rate for Direct-to-Consumer channel.', 'Growth rate for Retail channel.', 'Assumed average annual price increase across products.', 'Seasonal sales distribution (e.g., 0.2, 0.3, 0.3, 0.2).',
        # COGS Assumptions
        'Expected annual increase in raw material costs.', 'Average cost of ingredients per unit produced (baseline).', 'Expected annual increase in manufacturing labor wages.', 'Overhead applied based on direct labor costs.', 'Fee paid to contract manufacturers if applicable.', 'Cost of quality checks as a percentage of COGS.', 'Percentage of products lost during production.',
        # Operational Cost Assumptions
        'Expected annual increase in facility rent.', 'Expected annual increase in utility costs.', 'Cost of maintaining equipment.', 'Cost associated with holding inventory.', 'Annual cost for supply chain management software.', 'Annual budget for sustainability initiatives.',
        # Packaging & Logistics Assumptions
        'Expected annual increase in packaging material costs.', 'Annual cost per square foot for warehouse space.', 'Average cost to transport one unit of product.', 'Cost to process and ship each customer order.', 'Cost of final delivery to customer (esp. for DTC).', 'Percentage of sales value returned by customers.', 'Annual cost for logistics and tracking software.',
        # SG&A Assumptions
        'Projected growth in the number of sales personnel.', 'Average commission paid to sales staff.', 'Budget allocated to marketing and advertising.', 'Projected growth in administrative staff.', 'Expected annual increase in office rent.', 'Annual IT support cost per employee.', 'Estimated annual fees for legal and accounting services.', 'Budget for travel and entertainment expenses.',
        # R&D Assumptions
        'Budget allocated to research and development.', 'Estimated cost to bring a new product to market.', 'Cost associated with conducting clinical trials.', 'Cost to file for patent protection.', 'Annual cost for meeting regulatory requirements.', 'Annual costs for R&D collaborations.',
        # HR Assumptions
        'Average annual salary increase for employees.', 'Cost of employee benefits (health, retirement, etc.).', 'Annual budget for employee training.', 'Average cost to recruit a new employee.', 'Rate at which employees leave the company.', 'Cost of using temporary staff during peak times.',
        # Capex Assumptions
        'Budget for capital expenditures.', 'Expected lifespan of major equipment.', 'Annual budget for improving facilities.', 'Frequency of major technology system upgrades.', 'Method used for depreciation (e.g., Straight-Line).', 'Portion of Capex dedicated to maintaining existing assets.',
        # Working Capital Assumptions
        'Average days raw materials are held.', 'Average days finished goods are held before sale.', 'Average days it takes to collect payment from customers.', 'Average days it takes to pay suppliers.',
        # Group Assumptions
        'Markup applied to sales between group entities.', 'Methodology for setting prices on intercompany transactions.', 'How shared service costs are allocated (e.g., headcount, revenue).', 'How corporate overhead costs are distributed.', 'Manual adjustments needed for consolidation.',
        # Financial Assumptions
        'Interest rate applied to company debt.', 'Assumed exchange rate for key foreign currency.', 'Annual fees paid for banking services.',
        # Sustainability & Compliance Assumptions
        'Annual cost for maintaining certifications (organic, cruelty-free).', 'Estimated potential costs for environmental non-compliance.', 'Annual fees for financial and compliance audits.', 'Annual budget for corporate social responsibility programs.', 'Annual cost for auditing supply chain partners.',
    ]
}

assumptions_df = pd.DataFrame(assumptions_data)

# Create a new Excel file with pandas
with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
    # Write the DataFrame to the Excel file
    assumptions_df.to_excel(writer, sheet_name='Assumptions', index=False, startrow=1)
    
    # Access the workbook and sheet
    workbook = writer.book
    worksheet = writer.sheets['Assumptions']
    
    # --- Formatting --- 
    # Define styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    category_font = Font(bold=True)
    category_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    input_fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid") # Yellow fill for input cells
    
    # Add title
    worksheet['A1'] = 'Financial Model Assumptions'
    worksheet['A1'].font = Font(bold=True, size=14)
    worksheet.merge_cells('A1:H1') # Merge across all columns used
    worksheet['A1'].alignment = Alignment(horizontal='center')
    
    # Format header row (Row 2)
    for col_idx, value in enumerate(assumptions_df.columns, 1):
        cell = worksheet.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    
    # Format data rows (Starting from Row 3)
    last_category = None
    for row_idx in range(3, worksheet.max_row + 1):
        category_cell = worksheet.cell(row=row_idx, column=1)
        assumption_cell = worksheet.cell(row=row_idx, column=2)
        notes_cell = worksheet.cell(row=row_idx, column=8) # Notes column
    
        # Apply border to all cells in the row
        for col_idx in range(1, worksheet.max_column + 1):
            worksheet.cell(row=row_idx, column=col_idx).border = thin_border
    
        # Format Category column
        if category_cell.value and category_cell.value != last_category:
            last_category = category_cell.value
            category_cell.font = category_font
            category_cell.fill = category_fill
            category_cell.alignment = Alignment(vertical='center', horizontal='left', wrap_text=True)
    
        # Align Assumption and Notes columns
        assumption_cell.alignment = Alignment(vertical='center', horizontal='left', wrap_text=True)
        notes_cell.alignment = Alignment(vertical='center', horizontal='left', wrap_text=True)
    
        # Highlight input cells (Years 1-5)
        for col_idx in range(3, 8): # Columns C to G (Year 1 to Year 5)
            input_cell = worksheet.cell(row=row_idx, column=col_idx)
            input_cell.fill = input_fill
            input_cell.alignment = Alignment(horizontal='right') # Align numbers to the right
            # Add number formatting (optional, e.g., percentage)
            if assumption_cell.value and '%' in assumption_cell.value:
                input_cell.number_format = '0.00%'
            elif assumption_cell.value and 'Rate' in assumption_cell.value:
                 input_cell.number_format = '0.00%'
            elif assumption_cell.value and 'Days' in assumption_cell.value:
                 input_cell.number_format = '0'
            elif assumption_cell.value and ('Cost' in assumption_cell.value or 'Budget' in assumption_cell.value or 'Fees' in assumption_cell.value):
                 input_cell.number_format = '#,##0'
            elif assumption_cell.value and 'Year' in assumption_cell.value:
                 input_cell.number_format = '0'
    
    # Adjust column widths
    worksheet.column_dimensions['A'].width = 25 # Category
    worksheet.column_dimensions['B'].width = 40 # Assumption
    worksheet.column_dimensions['C'].width = 15 # Year 1
    worksheet.column_dimensions['D'].width = 15 # Year 2
    worksheet.column_dimensions['E'].width = 15 # Year 3
    worksheet.column_dimensions['F'].width = 15 # Year 4
    worksheet.column_dimensions['G'].width = 15 # Year 5
    worksheet.column_dimensions['H'].width = 50 # Notes/Drivers

print(f"Assumptions sheet created/updated in {excel_file}")
