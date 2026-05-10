from __future__ import annotations

import openpyxl
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from financial_model_config import load_existing_workbook, parse_output_args, resolve_output_path, save_workbook

TRIAD_ALLOCATIONS = [
    ('Triad 1 · Market Intelligence', 7020000),
    ('Triad 2 · Resource Management', 9600000),
    ('Triad 3 · Innovation', 7020000),
    ('Triad 4 · Operations', 6800000),
    ('Triad 5 · Customer Interface', 3900000),
    ('Triad 6 · Financial Optimization', 10300000),
    ('Triad 7 · Organizational Development', 5200000),
]

GROUP_METRIC_ROW = {
    'Revenue': 91,
    'Cost of Goods Sold (COGS)': 92,
    'Gross Profit': 93,
    'Gross Margin %': 94,
    'Operational Costs': 95,
    'SG&A Expenses': 96,
    'EBITDA': 97,
    'EBITDA Margin %': 98,
    'Net Income': 104,
    'Net Profit Margin %': 105,
}


def build_sheet(workbook: openpyxl.Workbook) -> None:
    if 'Dashboard' in workbook.sheetnames:
        dashboard_index = workbook.sheetnames.index('Dashboard')
        workbook.remove(workbook.worksheets[dashboard_index])

    dashboard = workbook.create_sheet('Dashboard', 0)

    title_font = Font(bold=True, size=16)
    section_font = Font(bold=True, size=13)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    subheader_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    dashboard['A1'] = 'Skin Care Supply Chain Financial Model Dashboard'
    dashboard['A1'].font = title_font
    dashboard['A1'].alignment = Alignment(horizontal='center')
    dashboard.merge_cells('A1:H1')

    dashboard['A3'] = 'Key Financial Metrics'
    dashboard['A3'].font = section_font
    dashboard.merge_cells('A3:H3')
    dashboard['A3'].alignment = Alignment(horizontal='center')

    dashboard['A4'] = 'Metric'
    dashboard['A4'].font = header_font
    dashboard['A4'].fill = header_fill
    dashboard['A4'].border = thin_border
    year_columns = ['Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5']
    for offset, year in enumerate(year_columns, start=2):
        cell = dashboard.cell(row=4, column=offset)
        cell.value = year
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    metrics = [
        'Revenue',
        'Gross Profit',
        'Gross Margin %',
        'EBITDA',
        'EBITDA Margin %',
        'Net Income',
        'Net Profit Margin %',
        'Operational Costs',
        'SG&A Expenses',
    ]

    for row_offset, metric in enumerate(metrics, start=5):
        name_cell = dashboard.cell(row=row_offset, column=1)
        name_cell.value = metric
        name_cell.font = Font(bold=True)
        name_cell.border = thin_border
        group_row = GROUP_METRIC_ROW[metric]
        for year_offset in range(5):
            data_cell = dashboard.cell(row=row_offset, column=year_offset + 2)
            group_col = chr(ord('B') + year_offset)
            data_cell.value = f"='Group Model'!{group_col}{group_row}"
            data_cell.border = thin_border
            data_cell.alignment = Alignment(horizontal='right')
            if '%' in metric:
                data_cell.number_format = '0.0%'
            else:
                data_cell.number_format = '#,##0'

    ratio_start_row = 16
    dashboard[f'A{ratio_start_row}'] = 'Key Ratios'
    dashboard[f'A{ratio_start_row}'].font = section_font
    dashboard.merge_cells(start_row=ratio_start_row, start_column=1, end_row=ratio_start_row, end_column=4)
    dashboard[f'A{ratio_start_row}'].alignment = Alignment(horizontal='center')

    ratio_rows = [
        ('Gross Margin %', 0.0),
        ('EBITDA Margin %', 0.0),
        ('R&D % of Revenue', 0.06),
        ('Headcount Growth Proxy', 0.05),
    ]
    header_row = ratio_start_row + 1
    for column, header in enumerate(['Ratio', 'Year 1', 'Year 3', 'Year 5'], start=1):
        cell = dashboard.cell(row=header_row, column=column)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for offset, (label, fallback) in enumerate(ratio_rows, start=header_row + 1):
        label_cell = dashboard.cell(row=offset, column=1)
        label_cell.value = label
        label_cell.font = Font(bold=True)
        label_cell.border = thin_border
        for col_index, group_col in enumerate(['B', 'D', 'F'], start=2):
            value_cell = dashboard.cell(row=offset, column=col_index)
            if label in GROUP_METRIC_ROW:
                value_cell.value = f"='Group Model'!{group_col}{GROUP_METRIC_ROW[label]}"
            else:
                value_cell.value = fallback
            value_cell.number_format = '0.0%'
            value_cell.border = thin_border

    triad_start_row = 24
    dashboard[f'A{triad_start_row}'] = 'Triad Allocation Summary'
    dashboard[f'A{triad_start_row}'].font = section_font
    dashboard.merge_cells(start_row=triad_start_row, start_column=1, end_row=triad_start_row, end_column=4)
    dashboard[f'A{triad_start_row}'].alignment = Alignment(horizontal='center')

    for column, header in enumerate(['Triad', 'Allocation', '% of Total'], start=1):
        cell = dashboard.cell(row=triad_start_row + 1, column=column)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    total_allocation = sum(value for _, value in TRIAD_ALLOCATIONS)
    for row_offset, (triad_name, amount) in enumerate(TRIAD_ALLOCATIONS, start=triad_start_row + 2):
        dashboard.cell(row=row_offset, column=1).value = triad_name
        dashboard.cell(row=row_offset, column=2).value = amount
        dashboard.cell(row=row_offset, column=2).number_format = '#,##0'
        dashboard.cell(row=row_offset, column=3).value = amount / total_allocation
        dashboard.cell(row=row_offset, column=3).number_format = '0.0%'
        for col in range(1, 4):
            dashboard.cell(row=row_offset, column=col).border = thin_border

    revenue_chart = BarChart()
    revenue_chart.title = 'Consolidated Revenue by Year'
    revenue_chart.y_axis.title = 'Revenue'
    revenue_chart.x_axis.title = 'Year'
    revenue_data = Reference(dashboard, min_col=2, min_row=5, max_col=6, max_row=5)
    revenue_categories = Reference(dashboard, min_col=2, min_row=4, max_col=6, max_row=4)
    revenue_chart.add_data(revenue_data, titles_from_data=False, from_rows=True)
    revenue_chart.set_categories(revenue_categories)
    revenue_chart.height = 8
    revenue_chart.width = 14
    dashboard.add_chart(revenue_chart, 'E4')

    cost_chart = PieChart()
    cost_chart.title = 'Triad Cost Allocation'
    cost_data = Reference(dashboard, min_col=2, min_row=triad_start_row + 2, max_row=triad_start_row + 8)
    cost_labels = Reference(dashboard, min_col=1, min_row=triad_start_row + 2, max_row=triad_start_row + 8)
    cost_chart.add_data(cost_data, titles_from_data=False)
    cost_chart.set_categories(cost_labels)
    cost_chart.height = 10
    cost_chart.width = 12
    dashboard.add_chart(cost_chart, 'E18')

    for row in range(5, 14):
        for column in range(1, 7):
            dashboard.cell(row=row, column=column).border = thin_border

    for row in range(1, dashboard.max_row + 1):
        for col in range(1, dashboard.max_column + 1):
            cell = dashboard.cell(row=row, column=col)
            existing_alignment = cell.alignment
            cell.alignment = Alignment(
                horizontal=existing_alignment.horizontal,
                vertical='center',
                wrap_text=existing_alignment.wrap_text,
            )

    dashboard.column_dimensions['A'].width = 28
    dashboard.column_dimensions['B'].width = 14
    dashboard.column_dimensions['C'].width = 14
    dashboard.column_dimensions['D'].width = 14
    dashboard.column_dimensions['E'].width = 16
    dashboard.column_dimensions['F'].width = 16
    dashboard.column_dimensions['G'].width = 16
    dashboard.column_dimensions['H'].width = 16


def main(output_path: str | None = None):
    excel_file = resolve_output_path(output_path)
    workbook = load_existing_workbook(excel_file, ['Assumptions', 'Revenue', 'COGS', 'Operational Costs', 'SG&A', 'Group Model'])
    build_sheet(workbook)
    save_workbook(workbook, excel_file)
    print(f'Dashboard sheet created in {excel_file}')
    return excel_file


if __name__ == '__main__':
    args = parse_output_args('Create the dashboard sheet for the SkinTwin financial model workbook.')
    main(args.output)
