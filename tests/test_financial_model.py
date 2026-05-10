from __future__ import annotations

from pathlib import Path
import sys

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from run_all import main as run_all_main


def build_workbook(tmp_path: Path):
    output_path = tmp_path / 'financial-model.xlsx'
    run_all_main(str(output_path))
    return load_workbook(output_path, data_only=False), output_path


def test_run_all_creates_expected_sheets(tmp_path: Path):
    workbook, output_path = build_workbook(tmp_path)

    assert output_path.exists()
    assert workbook.sheetnames[:3] == ['Dashboard', 'Assumptions', 'Group Model']
    assert {'Revenue', 'COGS', 'Operational Costs', 'SG&A'}.issubset(set(workbook.sheetnames))



def test_dashboard_contains_numeric_allocations_and_group_links(tmp_path: Path):
    workbook, _ = build_workbook(tmp_path)
    dashboard = workbook['Dashboard']

    assert dashboard['B5'].value == "='Group Model'!B91"
    assert dashboard['C5'].value == "='Group Model'!C91"
    assert isinstance(dashboard['B26'].value, (int, float))
    assert dashboard['B26'].value > 0
    assert dashboard['C26'].number_format == '0.0%'



def test_group_model_consolidated_formulas_reference_entity_rows(tmp_path: Path):
    workbook, _ = build_workbook(tmp_path)
    group_model = workbook['Group Model']

    assert group_model['B91'].value == '=B6+B22+B38+B54+B70'
    assert group_model['C91'].value == '=C6+C22+C38+C54+C70'
    assert group_model['B104'].value == '=B19+B35+B51+B67+B83'
    assert group_model['C104'].value == '=C19+C35+C51+C67+C83'
